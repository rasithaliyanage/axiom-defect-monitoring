"""Reconstruct the supplied Task 4 PDF's layout as an editable board guide."""
from pathlib import Path
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / '.tmp-doc-tools'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

STEM = ROOT / 'BOARD_DEFECT_INSPECTION_TASK_4_BEGINNER_GUIDE_ENTERPRISE_STANDARD'


def inline(p, text):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        r = p.add_run(part[2:-2] if part.startswith('**') else part)
        r.bold = part.startswith('**')


def main():
    content = STEM.with_suffix('.md').read_text(encoding='utf-8')
    pages = content.split('<!-- PAGE -->')
    assert len(pages) == 14
    numbers = re.findall(r'^## (\d+)\.', content, re.M)
    assert list(map(int, numbers)) == list(range(1, 18))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    # Source PDF text transform: left 57.6 pt, top 50.4 pt; US Letter.
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    n = doc.styles['Normal']
    n.font.name, n.font.size = 'Arial', Pt(10.5)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.05
    for name, size in [('Heading 1', 16), ('Heading 2', 13)]:
        s = doc.styles[name]
        s.font.name, s.font.size = 'Arial', Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(9)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    for index, page in enumerate(pages):
        first = True
        code = None
        for line in page.strip().splitlines():
            if line.startswith('```'):
                code = None if code else ('prompt' if line == '```prompt' else 'diagram')
                continue
            if not line.strip() and not code:
                continue
            if code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(10.2 if code == 'prompt' else 11)
                p.paragraph_format.widow_control = False
                r = p.add_run(line or ' ')
                r.font.name = 'Consolas'
                r.font.size = Pt(8.5 if code == 'prompt' else 9)
            elif line.startswith('# '):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(line[2:])
                # PDF embeds Play; use available Arial instead of Word's serif fallback.
                r.font.name, r.font.size = 'Arial', Pt(22)
            elif line.startswith('## '):
                p = doc.add_heading(line[3:], 1)
            elif line.startswith('### '):
                p = doc.add_heading(line[4:], 2)
            elif line.startswith('@ '):
                p = doc.add_paragraph(line[2:])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].bold = True
            elif line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                inline(p, line[2:])
            else:
                p = doc.add_paragraph()
                inline(p, line)
            if first and index:
                p.paragraph_format.page_break_before = True
            first = False
    doc.core_properties.title = 'Board Defect Inspection AI — Task 4 — Controlled UISpec Renderer'
    doc.core_properties.subject = 'Enterprise engineering guide and adapted implementation prompt'
    doc.core_properties.author = 'Board Defect Inspection Project'
    doc.core_properties.comments = 'Adapted from supplied 14-page PDF; planned work, not an implementation report.'
    doc.save(STEM.with_suffix('.docx'))
    with zipfile.ZipFile(STEM.with_suffix('.docx')) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml', '.rels')):
                ET.fromstring(z.read(name))
    prompt = '\n'.join(re.findall(r'```prompt\n(.*?)```', content, re.S))
    assert 'Do NOT commit. Do NOT push. Do NOT start Task 5.' in prompt
    assert all(x in prompt for x in ['Hero', 'Heading', 'Text', 'FeatureCard', 'FeatureGrid', 'CTA', 'Alert', 'InfoPanel', 'ProgressIndicator'])
    print('Created DOCX: 14 explicit page groups, all 17 numbered sections, complete adapted prompt.')
    print('Verified DOCX ZIP/XML integrity and required prompt coverage.')


if __name__ == '__main__':
    main()
